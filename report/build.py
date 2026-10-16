import sys, copy, json, re
sys.path.insert(0, "report")
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.text.paragraph import Paragraph
from PIL import Image
import content as C

TEMPLATE = "/mnt/user-data/uploads/1791898037280_CPPE_Mini_project_2_report_template.docx"
pages = json.load(open("report/pages.json")) if len(sys.argv) > 1 and sys.argv[1] == "final" else {}
doc = docx.Document(TEMPLATE); ps = list(doc.paragraphs)

# ---------- numbering (bullets + decimal) ----------
numel = doc.part.numbering_part.element
first_num = numel.find(qn("w:num"))
def add_abstract(aid, fmt, text):
    x = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="851" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>')
    first_num.addprevious(x) if first_num is not None else numel.append(x)
    numel.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{aid}"><w:abstractNumId w:val="{aid}"/></w:num>'))
add_abstract(90, "bullet", "•"); add_abstract(91, "decimal", "%1.")

# ---------- element builders ----------
def mk_p(jc="both", first=567, left=None, after=160, before=0, keep_next=False, numid=None, line=276):
    p = OxmlElement("w:p"); pPr = OxmlElement("w:pPr"); p.append(pPr)
    st = OxmlElement("w:pStyle"); st.set(qn("w:val"), "Normal1"); pPr.append(st)
    if keep_next: pPr.append(OxmlElement("w:keepNext"))
    if numid:
        n = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{numid}"/></w:numPr>'); pPr.append(n)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), str(before)); sp.set(qn("w:after"), str(after)); sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), "auto"); pPr.append(sp)
    ind = OxmlElement("w:ind")
    if numid: ind.set(qn("w:left"), "851"); ind.set(qn("w:hanging"), "360")
    else:
        if left is not None: ind.set(qn("w:left"), str(left))
        if first: ind.set(qn("w:firstLine"), str(first))
    pPr.append(ind)
    j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); pPr.append(j)
    return p
def add_run(p, text, bold=False, size=24, italic=False):
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    if bold: rPr.append(OxmlElement("w:b"))
    if italic: rPr.append(OxmlElement("w:i"))
    for t in ("w:sz", "w:szCs"):
        e = OxmlElement(t); e.set(qn("w:val"), str(size)); rPr.append(e)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); p.append(r); return r

def body_para(text):  return add_run(mk_p(), text) and None
def make_text_p(text): p = mk_p(); add_run(p, text); return p
def make_item_p(item, numid):
    p = mk_p(numid=numid, after=80)
    if isinstance(item, tuple): add_run(p, item[0], bold=True); add_run(p, item[1])
    else: add_run(p, item)
    return p
def make_img(path, w, cap):
    im = Image.open(path); ar = im.height/im.width; w = min(w, 3.9/ar) if ar > 1 else w
    holder = doc.add_paragraph(); pp = holder._p; pp.getparent().remove(pp)   # create via python-docx for relationship handling
    p = mk_p(jc="center", first=0, after=60, before=100, keep_next=True)
    para = Paragraph(p, holder._parent); para.add_run().add_picture(path, width=Inches(w))
    cap_p = mk_p(jc="center", first=0, after=220); add_run(cap_p, cap, size=22)
    return [p, cap_p]
def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, sz in kw.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000"); b.append(e)
def make_table(cap, head, rows, widths, align):
    cap_p = mk_p(jc="center", first=0, after=100, before=120, keep_next=True); add_run(cap_p, cap, size=22)
    tbl = doc.add_table(rows=1+len(rows), cols=len(head)); t = tbl._tbl
    tblPr = t.tblPr
    for e in list(tblPr): tblPr.remove(e)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{sum(widths)}" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
    tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="12" w:color="000000"/><w:bottom w:val="single" w:sz="12" w:color="000000"/></w:tblBorders>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    tblPr.append(parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="50" w:type="dxa"/><w:left w:w="80" w:type="dxa"/><w:bottom w:w="50" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>'))
    grid = t.find(qn("w:tblGrid"))
    for g, w in zip(grid.findall(qn("w:gridCol")), widths): g.set(qn("w:w"), str(w))
    align = align or ["l"]*len(head)
    if len(align) < len(head): align = align + [align[-1]]*(len(head)-len(align))
    for ri, row in enumerate([head]+rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.width = docx.shared.Twips(widths[ci])
            p = cell.paragraphs[0]; p.style = doc.styles["Normal1"]
            pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0)
            p.alignment = {"l": 0, "c": 1}[align[ci] if ri else ("l" if align[ci] == "l" else "c")]
            r = p.add_run(str(val)); r.font.size = Pt(10.5); r.bold = (ri == 0)
            if ri == 0: set_cell_border(cell, bottom=8)
        if ri == 0:
            trPr = tbl.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit")); trPr.append(OxmlElement("w:tblHeader"))
        else: tbl.rows[ri]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for ri2, row2 in enumerate(tbl.rows):
        if ri2 < len(tbl.rows)-1:
            for c2 in row2.cells:
                for p2 in c2.paragraphs: p2.paragraph_format.keep_with_next = True
    t.getparent().remove(t)
    spacer = mk_p(first=0, after=120); add_run(spacer, "", size=8)
    return [cap_p, t, spacer]

def insert(after_el, items):
    cur = after_el
    def put(el):
        nonlocal cur; cur.addnext(el); cur = el
    for n_i, it in enumerate(items):
        k = it[0]
        if k == "p":
            el = make_text_p(it[1])
            if n_i+1 < len(items) and items[n_i+1][0] in ("img", "tbl"): el.find(qn("w:pPr")).insert(1, OxmlElement("w:keepNext"))
            put(el)
        elif k == "b":
            for x in it[1]: put(make_item_p(x, 90))
            sp = mk_p(first=0, after=60); add_run(sp, "", size=8); put(sp)
        elif k == "n":
            for x in it[1]: put(make_item_p(x, 91))
            sp = mk_p(first=0, after=60); add_run(sp, "", size=8); put(sp)
        elif k == "img":
            for el in make_img(it[1], it[2], it[3]): put(el)
        elif k == "tbl":
            for el in make_table(*it[1:]): put(el)
    return cur

# ---------- front matter ----------
def hl(run): run.font.highlight_color = WD_COLOR_INDEX.YELLOW
p0 = ps[0]; p0.runs[0].text = C.TITLE; p0.alignment = 1
pf = p0.paragraph_format; pf.left_indent = Inches(0.4); pf.right_indent = Inches(0.4)
r3 = ps[3].runs; r3[0].text = "[YOUR NAME]"; hl(r3[0])
nr = copy.deepcopy(r3[-1]._r); r3[-1]._r.addnext(nr)
from docx.text.run import Run
rr = Run(nr, ps[3]); rr.text = "[REG NO]"; hl(rr)
ps[8].runs[0].text = "[GUIDE NAME]"; hl(ps[8].runs[0])
r29 = ps[29].runs; r29[1].text = C.TITLE; r29[3].text = "[YOUR NAME]"; hl(r29[3]); r29[5].text = "REG NO"; hl(r29[5])
r36 = ps[36].runs
for i in (8, 9, 10, 11, 15, 16, 17, 18, 19): hl(r36[i])
r36[22].text = "Dr. R. SARAVANAN"
d44 = ps[44]; nr = copy.deepcopy(d44.runs[0]._r); d44.runs[0]._r.addnext(nr); rr = Run(nr, d44); rr.text = "[DD OCTOBER 2026]"; hl(rr)

# ---------- locate & clean ----------
keep_sect = {143, 159}
def is_empty(p): 
    x = p._p.xml
    return not p.text.strip() and "w:drawing" not in x and "w:pict" not in x
for i in list(range(89, 116)) + [121,122,124,125,127,128,133,134,136,137,139,140,141,142,147,148,149,151,152,154,155,157,158,163,164,166,167,168,169,170,172,173] + list(range(176, 202)) + [203]:
    if i in keep_sect: continue
    p = ps[i]
    if is_empty(p) or p.text.strip() == ".": p._p.getparent().remove(p._p)

# move section breaks from 143 / 159 to the last content paragraph
def take_sect(p):
    s = p._p.pPr.find(qn("w:sectPr")); p._p.pPr.remove(s); return s

cur = {}
def sec(idx, items): return insert(ps[idx]._p, items)
sec(87, C.INTRO); sec(88, C.PROB); sec(116, C.OBJ); sec(117, C.ARCH)
sec(120, C.M1_PROC); sec(123, C.M1_SHOT); sec(126, C.M1_ACT); sec(129, C.M1_RES)
sec(132, C.M2_PROC); sec(135, C.M2_SHOT); sec(138, C.M2_ACT)
s143 = take_sect(ps[143]); last = sec(143, C.M2_RES)
Paragraph(last, ps[143]._parent)._p.get_or_add_pPr().append(s143)
sec(146, C.M3_PROC); sec(150, C.M3_SHOT); sec(153, C.M3_ACT)
s159 = take_sect(ps[159]); last = sec(156, C.M3_RES)
Paragraph(last, ps[156]._parent)._p.get_or_add_pPr().append(s159)
ps[159]._p.getparent().remove(ps[159]._p)
sec(162, C.M4_PROC); sec(165, C.M4_SHOT); sec(171, C.M4_ACT); sec(174, C.M4_RES)
sec(175, C.CONC)
refs = []
for i, t in enumerate(C.REFS, 1):
    p = mk_p(first=0, left=567, after=100); p.find(qn("w:pPr")).find(qn("w:ind")).set(qn("w:hanging"), "567")
    add_run(p, f"[{i}]\t{t}", size=24); refs.append(p)
cur_el = ps[202]._p
for p in refs: cur_el.addnext(p); cur_el = p

# ---------- sections: footers / margins ----------
body = doc.element.body
sects = body.findall(".//" + qn("w:sectPr"))
RID_FOOT = "rId12"
for idx, s in enumerate(sects):
    if idx >= 2:   # intro onwards
        pm = s.find(qn("w:pgMar")); pm.set(qn("w:bottom"), "1100"); pm.set(qn("w:footer"), "500"); pm.set(qn("w:top"), "1300")
        if s.find(qn("w:footerReference")) is None:
            fr = OxmlElement("w:footerReference"); fr.set(qn("w:type"), "default"); fr.set(qn("r:id"), RID_FOOT)
            hr = s.find(qn("w:headerReference")); (hr.addnext(fr) if hr is not None else s.insert(0, fr))

# ---------- TOC page numbers ----------
toc_map = {62: "INTRODUCTION", 63: "PROBLEM STATEMENT", 64: "ARCHITECTURE DIAGRAM", 65: "MODULE 1", 70: "MODULE 2", 75: "MODULE 3", 80: "MODULE 4", 85: "CONCLUSION", 86: "REFERENCES"}
for i in range(66, 70): toc_map[i] = ps[i].text.split("\t")[1] + " " + ps[i].text.split("\t")[2]
for i in range(71, 75): toc_map[i] = ps[i].text.split("\t")[1] + " " + ps[i].text.split("\t")[2]
for i in range(76, 80): toc_map[i] = ps[i].text.split("\t")[1] + " " + ps[i].text.split("\t")[2]
for i in range(81, 85): toc_map[i] = ps[i].text.split("\t")[1] + " " + ps[i].text.split("\t")[2]
if pages:
    for i, key in toc_map.items():
        p = ps[i]; pg = str(pages[key])
        if i == 62:
            p.runs[7].text = pg
        else:
            nr = copy.deepcopy(p.runs[-1]._r); p.runs[-1]._r.addnext(nr); Run(nr, p).text = pg
json.dump({str(k): v for k, v in toc_map.items()}, open("report/toc_keys.json", "w"))
out = "report/SysPulse_Mini_Project_Report.docx" 
doc.save(out); print("saved", out)
